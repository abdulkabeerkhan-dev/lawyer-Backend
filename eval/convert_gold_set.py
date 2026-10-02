"""Convert the GoldSet sheet to JSON for the test harness.

Usage: python convert_gold_set.py gold_set_template.xlsx gold_set.json [--approved-only]
Multi-line cells become lists. Authorities are parsed from 'citation | needed for | min court'.
"""
import json, sys
from openpyxl import load_workbook

HEADER_ROW, FIRST = 4, 5
FIELDS = ["id","area","anchors","query","forum","date_sensitive","required_points","required_authorities",
          "forbidden_claims","expected_negative_findings","expected_forum_route","expected_limitation",
          "advice_risk","currency_check","sources","split","difficulty","author","reviewer","review_date","status","notes"]
LIST_FIELDS = {"anchors","required_points","forbidden_claims","expected_negative_findings"}

def lines(v):
    return [x.strip() for x in str(v or "").split("\n") if x.strip()]

def authority(line):
    parts = [p.strip() for p in line.split("|")]
    return {"citation": parts[0], "needed_for": parts[1] if len(parts) > 1 else "",
            "min_court": parts[2] if len(parts) > 2 else ""}

def main(src, dst, approved_only=False):
    ws = load_workbook(src, data_only=True)["GoldSet"]
    out = []
    for row in ws.iter_rows(min_row=FIRST, values_only=True):
        rec = dict(zip(FIELDS, row))
        if not rec.get("query"):
            continue
        if approved_only and rec.get("status") != "Approved":
            continue
        for f in LIST_FIELDS:
            rec[f] = lines(rec.get(f))
        rec["required_authorities"] = [authority(l) for l in lines(rec.get("required_authorities"))]
        rec["date_sensitive"] = (rec.get("date_sensitive") == "Yes")
        out.append(rec)
    with open(dst, "w", encoding="utf-8") as fh:
        json.dump(out, fh, indent=2, ensure_ascii=False, default=str)
    print(f"wrote {len(out)} cases to {dst}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], "--approved-only" in sys.argv)
