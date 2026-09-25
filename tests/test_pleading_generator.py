import unittest
import os
import sys
import asyncio
import io
from unittest.mock import AsyncMock, patch, MagicMock

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from prompts.pleading_generator import PLEADING_SYSTEM_PROMPT
from core.document_builder import generate_court_docx
import main
from docx import Document


class TestPleadingGenerator(unittest.TestCase):

    def test_stage_1_prompt_rules(self):
        self.assertIn("MANDATORY RULES:", PLEADING_SYSTEM_PROMPT)
        self.assertIn("SEPARATION OF FACT AND LAW", PLEADING_SYSTEM_PROMPT)
        self.assertIn("THE BRACKET RULE", PLEADING_SYSTEM_PROMPT)
        self.assertIn("**[INSERT DATE OF ORDER]**", PLEADING_SYSTEM_PROMPT)
        self.assertIn("ZERO PARAPHRASING OF LAW", PLEADING_SYSTEM_PROMPT)
        self.assertIn("STRUCTURE:", PLEADING_SYSTEM_PROMPT)
        self.assertIn("NO COMMENTARY:", PLEADING_SYSTEM_PROMPT)

    def test_stage_2_generate_court_docx(self):
        case_title = "Mst. Aisha Bibi v. Federation of Pakistan"
        court_name = "Lahore High Court, Lahore"
        pleading_text = (
            "1. That the petitioner is a law-abiding citizen of Pakistan.\n\n"
            "2. That the respondent passed the impugned order dated **[INSERT DATE OF ORDER]** without jurisdiction.\n\n"
            "3. As held by the august Supreme Court in 2024 SCMR 101:\n"
            "\"An administrative order without jurisdictional backing is void ab initio.\"\n\n"
            "PRAYER:\n"
            "It is most respectfully prayed that this petition be accepted."
        )

        buffer = generate_court_docx(case_title, court_name, pleading_text)
        self.assertIsInstance(buffer, io.BytesIO)
        docx_bytes = buffer.getvalue()
        self.assertGreater(len(docx_bytes), 1000)

        doc = Document(io.BytesIO(docx_bytes))
        section = doc.sections[0]
        self.assertAlmostEqual(section.left_margin.inches, 1.5, places=2)
        self.assertAlmostEqual(section.right_margin.inches, 1.0, places=2)
        self.assertAlmostEqual(section.top_margin.inches, 1.0, places=2)
        self.assertAlmostEqual(section.bottom_margin.inches, 1.0, places=2)

        text_content = "\n".join([p.text for p in doc.paragraphs])
        self.assertIn("IN THE LAHORE HIGH COURT, LAHORE", text_content)
        self.assertIn(case_title, text_content)
        self.assertIn("**[INSERT DATE OF ORDER]**", text_content)

    def test_stage_3_endpoint_empty_payload(self):
        with self.assertRaises(main.HTTPException) as ctx:
            asyncio.run(main.generate_pleading({"selected_case_ids": []}))
        self.assertEqual(ctx.exception.status_code, 400)

    def test_stage_3_endpoint_success_flow(self):
        mock_cases = [
            {
                "id": "01b7d342-b7fa-4b2b-9c62-1406e77a17e3",
                "case_id": "2024_CLC_1486",
                "neutral_citation": "2024 CLC 1486",
                "full_text": "The jurisdiction under Article 199 is discretionary and equitable. The petitioner must approach the court with clean hands."
            }
        ]

        mock_supabase_table = MagicMock()
        mock_select = MagicMock()
        mock_in = MagicMock()
        mock_execute = MagicMock(return_value=MagicMock(data=mock_cases))

        mock_in.execute = mock_execute
        mock_select.in_ = MagicMock(return_value=mock_in)
        mock_supabase_table.select = MagicMock(return_value=mock_select)

        mock_supabase = MagicMock()
        mock_supabase.table = MagicMock(return_value=mock_supabase_table)

        mock_response = MagicMock()
        mock_content_block = MagicMock()
        mock_content_block.text = (
            "IN THE HIGH COURT OF SINDH AT KARACHI\n\n"
            "1. The petitioner states that their account was frozen on **[INSERT DATE]**.\n\n"
            "2. As held in 2024 CLC 1486: 'The jurisdiction under Article 199 is discretionary...'\n\n"
            "PRAYER:\n"
            "It is prayed that the order be set aside."
        )
        mock_response.content = [mock_content_block]

        payload = {
            "facts": "Bank account frozen without SBP notice.",
            "court_name": "High Court of Sindh",
            "case_title": "ABC Corp v. State Bank",
            "selected_case_ids": ["01b7d342-b7fa-4b2b-9c62-1406e77a17e3"]
        }

        with patch.object(main, 'supabase', mock_supabase), \
             patch.object(main, 'safe_create_anthropic_message', AsyncMock(return_value=mock_response)):

            streaming_resp = asyncio.run(main.generate_pleading(payload))
            self.assertEqual(streaming_resp.media_type, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            self.assertIn("attachment; filename=draft_pleading.docx", streaming_resp.headers.get("content-disposition", ""))
            
            async def read_stream(resp):
                chunks = []
                async for c in resp.body_iterator:
                    chunks.append(c if isinstance(c, bytes) else c.encode("utf-8"))
                return b"".join(chunks)

            docx_bytes = asyncio.run(read_stream(streaming_resp))
            doc = Document(io.BytesIO(docx_bytes))
            content = "\n".join([p.text for p in doc.paragraphs])
            self.assertIn("IN THE HIGH COURT OF SINDH", content)
            self.assertIn("ABC Corp v. State Bank", content)
            self.assertIn("**[INSERT DATE]**", content)



if __name__ == '__main__':
    unittest.main()
