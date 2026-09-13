import unittest
import os
import sys
import io
import base64
import docx

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

class TestDocxDocExtraction(unittest.TestCase):

    def test_docx_extraction_with_octet_stream_mime(self):
        # Build synthetic DOCX document in memory
        doc = docx.Document()
        doc.add_paragraph("SUIT FOR PARTITION AND MESNE PROFITS")
        doc.add_paragraph("Plaint filed under Order 7 Rule 1 CPC.")
        
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()
        b64_docx = base64.b64encode(docx_bytes).decode("utf-8")

        import main
        text = main.extract_text_from_document_base64(b64_docx, "application/octet-stream")
        self.assertIn("SUIT FOR PARTITION", text)
        self.assertIn("Order 7 Rule 1 CPC", text)

if __name__ == '__main__':
    unittest.main()
